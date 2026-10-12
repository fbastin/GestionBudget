"""
Reconcile the app's database with external sources and report the gaps.

Two optional sources, described in a JSON configuration (see reconcile_config.example.json):
- a spreadsheet of debts (.xlsx, local or fetched with rclone): one block of rows, one debt
  per row, its name in one column and its current balance in another;
- a brokerage holdings report (.csv, the latest file matching a pattern in a directory, by
  default Wealthsimple's holdings-report-YYYY-MM-DD.csv): market value of each account,
  converted to the base currency.

By default nothing is written: the comparison goes to a Markdown report. With --apply, the
chosen balances are recorded as snapshots through the app's API (which must be running), so
that the app's own logic applies, as if they had been entered by hand.

    python reconcile.py                                   # report only
    python reconcile.py --apply accounts                  # record every account gap
    python reconcile.py --apply debts --only Visa Loan    # record these debts only

The configuration and the report hold personal names and amounts: keep them out of git.
"""
import argparse
import copy
import csv
import glob
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import urllib.request
from collections import defaultdict
from datetime import date, datetime, timedelta

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(BASE_DIR, 'reconcile_config.json')
EXAMPLE_CONFIG_PATH = os.path.join(BASE_DIR, 'reconcile_config.example.json')
REPORT_PATH = os.path.join(BASE_DIR, 'reconcile_report.md')
TOLERANCE = 0.01

DEFAULT_CONFIG = {
    'api_url': 'http://127.0.0.1:8000',
    'spreadsheet': {
        'path': '',            # a local .xlsx file...
        'rclone_remote': '',   # ...or one fetched with rclone, e.g. "gdrive:debts.xlsx"
        'sheet': '',           # default: the current year, e.g. "2026"
        'name_column': 'A',
        'balance_column': 'B',
        'first_row': 2,        # the block ends at the first row without a name
    },
    'holdings': {
        'directory': '',
        'pattern': 'holdings-report-*.csv',
        'base_currency': 'CAD',
        'fx_rates': {},        # e.g. {"USD": 1.37}; otherwise fetched from the Bank of Canada (base CAD only)
        'columns': {
            'account': 'Account Name',
            'account_type': 'Account Type',
            'classification': 'Account Classification',
            'symbol': 'Symbol',
            'market_value': 'Market Value',
            'currency': 'Market Value Currency',
        },
    },
    # Source name -> name in the app, when they differ (others match by name, ignoring case)
    'debts': {},
    'debts_ignore': [],
    # Account key ("account|type|classification", empty parts left out) -> asset name in the app
    'accounts': {},
    'accounts_ignore': [],
}


def normalize(name):
    return re.sub(r'\s+', ' ', str(name or '')).strip().casefold()


def merge(defaults, values):
    merged = copy.deepcopy(defaults)
    for key, value in values.items():
        if isinstance(value, dict) and isinstance(merged.get(key), dict):
            merged[key] = merge(merged[key], value)
        else:
            merged[key] = value
    return merged


def load_config(path):
    if not os.path.exists(path):
        shutil.copyfile(EXAMPLE_CONFIG_PATH, path)
        sys.exit(f'Configuration créée à partir de l\'exemple : {path}. La compléter, puis relancer.')
    with open(path, encoding='utf-8') as f:
        return merge(DEFAULT_CONFIG, json.load(f))


# -- Spreadsheet of debts ----------------------------------------------------------------

def spreadsheet_file(settings, workdir):
    """Local path of the spreadsheet and the date it was last modified, or (None, None)."""
    if settings['path']:
        return settings['path'], date.fromtimestamp(os.path.getmtime(settings['path']))
    if settings['rclone_remote']:
        remote = settings['rclone_remote']
        info = json.loads(subprocess.run(['rclone', 'lsjson', '--stat', remote], check=True,
                                         capture_output=True, text=True).stdout)
        local = os.path.join(workdir, os.path.basename(remote.split(':', 1)[-1]) or 'debts.xlsx')
        subprocess.run(['rclone', 'copyto', remote, local], check=True)
        return local, datetime.fromisoformat(info['ModTime'].replace('Z', '+00:00')).astimezone().date()
    return None, None


def read_debts(path, sheet, name_column='A', balance_column='B', first_row=2):
    """[(name, balance or None)] of the block, from first_row down to the first row without a name."""
    import openpyxl
    from openpyxl.utils import column_index_from_string

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in workbook.sheetnames:
        sys.exit(f'Pas de feuille « {sheet} » dans {path}.')
    name_index = column_index_from_string(name_column) - 1
    balance_index = column_index_from_string(balance_column) - 1
    debts = []
    for row in workbook[sheet].iter_rows(min_row=first_row, values_only=True):
        name = row[name_index] if name_index < len(row) else None
        if name is None or str(name).strip() == '':
            break
        balance = row[balance_index] if balance_index < len(row) else None
        debts.append((str(name).strip(), float(balance) if isinstance(balance, (int, float)) else None))
    return debts


# -- Holdings report ---------------------------------------------------------------------

def holdings_file(directory, pattern):
    """Latest report matching the pattern and its date (from its name, otherwise its modification date)."""
    files = sorted(glob.glob(os.path.join(directory, pattern)))
    if not files:
        sys.exit(f'Aucun fichier {pattern} dans {directory}.')
    path = files[-1]
    found = re.search(r'(\d{4}-\d{2}-\d{2})', os.path.basename(path))
    return path, date.fromisoformat(found.group(1)) if found else date.fromtimestamp(os.path.getmtime(path))


def bank_of_canada_rate(currency, day):
    """Rate of `currency` in CAD on the last business day up to `day`."""
    series = f'FX{currency}CAD'
    url = (f'https://www.bankofcanada.ca/valet/observations/{series}/json'
           f'?start_date={day - timedelta(days=10)}&end_date={day}')
    with urllib.request.urlopen(url, timeout=20) as response:
        observations = json.load(response)['observations']
    if not observations:
        sys.exit(f'Pas de taux {series} publié par la Banque du Canada autour du {day} : le fixer dans fx_rates.')
    return float(observations[-1][series]['v'])


def exchange_rates(currencies, base, fixed, day):
    """{currency: rate in the base currency} for every currency of the report."""
    rates = {base: 1.0}
    for currency in sorted(set(currencies) - {base}):
        if currency in fixed:
            rates[currency] = float(fixed[currency])
        elif base == 'CAD':
            rates[currency] = bank_of_canada_rate(currency, day)
        else:
            sys.exit(f'Taux {currency}/{base} inconnu : le fixer dans holdings.fx_rates.')
    return rates


def read_holdings(path, columns):
    """Rows of the report that hold a position (the closing "As of ..." line has no symbol)."""
    with open(path, encoding='utf-8-sig') as f:
        return [row for row in csv.DictReader(f)
                if (row.get(columns['account']) or '').strip() and (row.get(columns['symbol']) or '').strip()]


def account_key(row, columns):
    parts = [(row.get(columns[c]) or '').strip() for c in ('account', 'account_type', 'classification') if columns.get(c)]
    return '|'.join(p for p in parts if p)


def account_values(rows, columns, rates, base):
    """{account key: market value in the base currency} and {account key: number of positions}."""
    totals, positions = defaultdict(float), defaultdict(int)
    for row in rows:
        key = account_key(row, columns)
        currency = (row.get(columns['currency']) or '').strip() or base
        totals[key] += float(row[columns['market_value']] or 0) * rates[currency]
        positions[key] += 1
    return totals, positions


# -- Database ----------------------------------------------------------------------------

def database_path():
    return os.environ.get('FINANCE_DB_PATH', os.path.join(BASE_DIR, 'finance.db'))


def latest_snapshots(table):
    """{normalized name: (id, name, date of the last snapshot, its amount)} of liabilities or assets."""
    column = {'liabilities': 'liability_id', 'assets': 'asset_id'}[table]
    db = sqlite3.connect(f'file:{database_path()}?mode=ro', uri=True)
    items = {}
    for item_id, name in db.execute(f'SELECT id, name FROM {table}'):
        last = db.execute(f'SELECT date, amount FROM snapshots WHERE {column} = ? ORDER BY date DESC LIMIT 1',
                          (item_id,)).fetchone()
        items[normalize(name)] = (item_id, name, *(last or (None, None)))
    db.close()
    return items


def investment_counts():
    db = sqlite3.connect(f'file:{database_path()}?mode=ro', uri=True)
    counts = dict(db.execute('SELECT asset_id, COUNT(*) FROM investments GROUP BY asset_id'))
    db.close()
    return counts


# -- Comparison and report ---------------------------------------------------------------

def compare(sources, mapping, ignore, items):
    """Rows (source name, app item or None, source amount, gap) and the app items left unmatched."""
    rows, matched = [], set()
    for name, amount in sources:
        if name in ignore:
            continue
        item = items.get(normalize(mapping.get(name, name)))
        if item:
            matched.add(item[0])
        gap = None if item is None or amount is None or item[3] is None else amount - item[3]
        rows.append((name, item, amount, gap))
    return rows, [item for item in items.values() if item[0] not in matched]


def has_gap(row):
    _, item, amount, gap = row
    return item is not None and amount is not None and (gap is None or abs(gap) >= TOLERANCE)


def money(value, currency):
    return '—' if value is None else f'{value:,.2f} {currency}'.replace(',', ' ')


def shown(name):
    # Account keys contain "|", which would split the Markdown table
    return name.replace('|', ' · ')


def gap_table(rows, source_label, currency):
    lines = [f'| {source_label} | Dans l\'app | Source | Dernier relevé de l\'app | Écart |', '|---|---|---|---|---|']
    for name, item, amount, gap in rows:
        if item is None:
            lines.append(f'| {shown(name)} | **aucune correspondance** | {money(amount, currency)} | — | — |')
            continue
        last = f'{money(item[3], currency)} au {item[2]}' if item[2] else 'aucun relevé'
        flag = ' ⚠' if gap is not None and abs(gap) >= TOLERANCE else ''
        lines.append(f'| {shown(name)} | {item[1]} | {money(amount, currency)} | {last} | {money(gap, currency)}{flag} |')
    return lines


def unmatched_list(items, currency):
    return [f'- {item[1]} (dernier relevé : {money(item[3], currency)} au {item[2]})' for item in items] or ['- aucun']


# -- Applying ----------------------------------------------------------------------------

def post_snapshot(api_url, payload):
    request = urllib.request.Request(f'{api_url}/snapshots', data=json.dumps(payload).encode(),
                                     headers={'Content-Type': 'application/json'}, method='POST')
    with urllib.request.urlopen(request, timeout=30) as response:
        return json.load(response)


def apply(rows, kind, day, only, send):
    """Record the source amount as a snapshot dated `day` for every row with a gap (or only the given names)."""
    key = {'debts': 'liability_id', 'accounts': 'asset_id'}[kind]
    recorded = []
    for row in rows:
        name, item, amount, _ = row
        if not has_gap(row) or (only and name not in only):
            continue
        send({'date': day.isoformat(), 'amount': round(amount, 2), key: item[0]})
        recorded.append(item[1])
    return recorded


# -- Command line ------------------------------------------------------------------------

def main(argv=None, send=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--config', default=CONFIG_PATH, help='JSON configuration (default: %(default)s)')
    parser.add_argument('--report', default=REPORT_PATH, help='Markdown report to write (default: %(default)s)')
    parser.add_argument('--apply', choices=['debts', 'accounts'], help='record the source amounts that differ')
    parser.add_argument('--only', nargs='+', default=[], help='restrict --apply to these source names')
    parser.add_argument('--sheet', help='spreadsheet sheet to read (default: configuration, else the current year)')
    args = parser.parse_args(argv)
    config = load_config(args.config)
    send = send or (lambda payload: post_snapshot(config['api_url'], payload))
    currency = config['holdings']['base_currency']
    lines = [f'# Rapprochement du {date.today()}', '']
    debt_rows = account_rows = None

    spreadsheet = config['spreadsheet']
    with tempfile.TemporaryDirectory() as workdir:
        path, modified = spreadsheet_file(spreadsheet, workdir)
        if path:
            sheet = args.sheet or spreadsheet['sheet'] or str(date.today().year)
            debts = read_debts(path, sheet, spreadsheet['name_column'], spreadsheet['balance_column'],
                               int(spreadsheet['first_row']))
    if path:
        debt_rows, debts_alone = compare(debts, config['debts'], config['debts_ignore'], latest_snapshots('liabilities'))
        lines += [f'- Classeur : `{spreadsheet["path"] or spreadsheet["rclone_remote"]}`, feuille {sheet}, '
                  f'modifié le {modified}.']

    holdings = config['holdings']
    if holdings['directory']:
        report, report_date = holdings_file(holdings['directory'], holdings['pattern'])
        columns = holdings['columns']
        rows = read_holdings(report, columns)
        rates = exchange_rates([(r.get(columns['currency']) or '').strip() or currency for r in rows],
                               currency, holdings['fx_rates'], report_date)
        totals, positions = account_values(rows, columns, rates, currency)
        account_rows, assets_alone = compare(sorted(totals.items()), config['accounts'], config['accounts_ignore'],
                                             latest_snapshots('assets'))
        others = ', '.join(f'{c}/{currency} {r}' for c, r in rates.items() if c != currency) or 'aucun change'
        lines += [f'- Relevé de positions : `{os.path.basename(report)}` ({others}).']

    if debt_rows is None and account_rows is None:
        sys.exit('Aucune source configurée : renseigner spreadsheet ou holdings.directory.')
    lines += [f'- Écart = source − dernier relevé de l\'app ; ⚠ au-delà de {TOLERANCE} {currency}.', '']

    if debt_rows is not None:
        lines += ['## Dettes', '', *gap_table(debt_rows, 'Classeur', currency), '',
                  '### Dettes de l\'app absentes du classeur', '', *unmatched_list(debts_alone, currency), '']
    if account_rows is not None:
        counts = investment_counts()
        lines += ['## Comptes (valeur marchande)', '', *gap_table(account_rows, 'Compte du relevé', currency), '',
                  '### Actifs de l\'app sans compte dans le relevé', '', '(normal pour les comptes d\'autres institutions)',
                  '', *unmatched_list(assets_alone, currency), '',
                  '## Positions', '', '| Compte du relevé | Positions au relevé | Placements dans l\'app |', '|---|---|---|']
        lines += [f'| {shown(name)} | {positions[name]} | {counts.get(item[0], 0) if item else "—"} |'
                  for name, item, _, _ in account_rows]

    with open(args.report, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')
    print(f'Rapport : {args.report}')
    for label, rows in (('Dettes', debt_rows), ('Comptes', account_rows)):
        if rows is not None:
            print(f'{label} : {len(rows)}, dont {sum(map(has_gap, rows))} avec écart et '
                  f'{sum(1 for r in rows if r[1] is None)} sans correspondance.')

    if args.apply:
        rows, day = (debt_rows, modified) if args.apply == 'debts' else (account_rows, report_date if account_rows is not None else None)
        if rows is None:
            sys.exit(f'--apply {args.apply} : cette source n\'est pas configurée.')
        recorded = apply(rows, args.apply, day, set(args.only), send)
        for name in recorded:
            print(f'  enregistré : {name} au {day}')
        print(f'{len(recorded)} relevé(s) enregistré(s).')


if __name__ == '__main__':
    main()
